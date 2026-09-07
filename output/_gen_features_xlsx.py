# -*- coding: utf-8 -*-
"""解析 output/features.md 中的功能清单表格，生成 Excel 文件。"""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = r"d:\AIProject\capPlatform\output\features.md"
DST = r"d:\AIProject\capPlatform\output\features.xlsx"

rows = []
with open(SRC, encoding="utf-8") as f:
    for line in f:
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if cells and re.fullmatch(r":?-{2,}:?", cells[0]) and all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
            continue
        if cells and cells[0] == "一级菜单":
            continue
        if len(cells) >= 5:
            rows.append(cells[:5])

wb = Workbook()
ws = wb.active
ws.title = "功能清单"

headers = ["一级菜单", "二级菜单", "功能点", "功能点说明", "所属角色"]
ws.append(headers)

header_fill = PatternFill("solid", fgColor="4472C4")
header_font = Font(bold=True, color="FFFFFF", size=11)
thin = Side(style="thin", color="BFBFBF")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(vertical="center", wrap_text=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

for col in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center
    cell.border = border

for r in rows:
    ws.append(r)

for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=5):
    for cell in row:
        cell.border = border
        if cell.column in (1, 2, 3, 5):
            cell.alignment = center
        else:
            cell.alignment = wrap

widths = [16, 24, 13, 70, 16]
for i, w in enumerate(widths, start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 22
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:E{ws.max_row}"

wb.save(DST)
print(f"OK rows={len(rows)} -> {DST}")
